# -*- coding: utf-8 -*-
"""把 作品集报告.md 和 业务说明文档.md 转成 docx（支持标题/表格/图片/列表/粗体）"""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.abspath(__file__))

def add_rich(par, text):
    """处理 **粗体** 行内标记"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for p in parts:
        if p.startswith("**") and p.endswith("**"):
            r = par.add_run(p[2:-2]); r.bold = True
        elif p:
            par.add_run(p)

def md_to_docx(md_path, docx_path, img_base):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"; style.font.size = Pt(10.5)
    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=min(level, 4))
            i += 1; continue
        if ln.strip().startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", ln.strip())
            if m:
                p = os.path.normpath(os.path.join(img_base, m.group(2)))
                if os.path.exists(p):
                    doc.add_picture(p, width=Inches(6.2))
                    cap = doc.add_paragraph(m.group(1))
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.runs[0].font.size = Pt(9)
            i += 1; continue
        if ln.strip().startswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i].strip()); i += 1
            rows = []
            for tl in tbl_lines:
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if all(re.match(r"^:?-+:?$", c) for c in cells):
                    continue
                rows.append(cells)
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        c = t.cell(ri, ci)
                        c.text = ""
                        add_rich(c.paragraphs[0], cell)
                        if ri == 0:
                            for r in c.paragraphs[0].runs:
                                r.bold = True
                        for r in c.paragraphs[0].runs:
                            r.font.size = Pt(9)
            continue
        if ln.strip().startswith((">", "- ")) and not ln.strip().startswith("---"):
            txt = ln.strip().lstrip(">").lstrip("- ").strip()
            p = doc.add_paragraph(style="List Bullet" if ln.strip().startswith("- ") else None)
            add_rich(p, txt)
            i += 1; continue
        if ln.strip() == "---":
            i += 1; continue
        p = doc.add_paragraph()
        add_rich(p, ln.strip())
        i += 1
    doc.save(docx_path)
    print("saved", docx_path)

md_to_docx(os.path.join(BASE, "reports", "作品集报告.md"),
           os.path.join(BASE, "reports", "作品集报告.docx"),
           os.path.join(BASE, "reports"))
md_to_docx(os.path.join(BASE, "docs", "业务说明文档.md"),
           os.path.join(BASE, "docs", "业务说明文档.docx"),
           os.path.join(BASE, "docs"))
